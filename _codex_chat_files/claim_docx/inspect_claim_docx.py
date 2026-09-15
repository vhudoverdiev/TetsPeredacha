from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def q(name: str) -> str:
    return f"{W}{name}"


path = Path(r"C:\Users\Владимир\Desktop\Пример притензии.docx")
doc = Document(path)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
for index, paragraph in enumerate(doc.paragraphs[:120]):
    fmt = paragraph.paragraph_format
    print(
        "P",
        index,
        repr(paragraph.text),
        "align",
        paragraph.alignment,
        "first",
        fmt.first_line_indent,
        "left",
        fmt.left_indent,
        "after",
        fmt.space_after,
        "before",
        fmt.space_before,
        "style",
        paragraph.style.name,
    )
for table_index, table in enumerate(doc.tables):
    print("TABLE", table_index, "rows", len(table.rows), "cols", len(table.columns))
    for row in table.rows[:6]:
        print([cell.text for cell in row.cells])
    print("widths", [cell.width for cell in table.rows[0].cells])

with ZipFile(path) as archive:
    document_xml = archive.read("word/document.xml")
root = ET.fromstring(document_xml)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
sect = root.find(".//w:sectPr", ns)
print("sectPr")
for child in list(sect):
    print(ET.tostring(child, encoding="unicode"))
for table_index, table in enumerate(root.findall(".//w:tbl", ns)):
    print("XML_TABLE", table_index)
    tbl_pr = table.find("w:tblPr", ns)
    print(ET.tostring(tbl_pr, encoding="unicode") if tbl_pr is not None else None)
    grid = table.find("w:tblGrid", ns)
    print(ET.tostring(grid, encoding="unicode") if grid is not None else None)
    first_row = table.find("w:tr", ns)
    print(ET.tostring(first_row, encoding="unicode")[:3000] if first_row is not None else None)
for index, para in enumerate(root.findall(".//w:p", ns)[:140]):
    texts = "".join(t.text or "" for t in para.findall(".//w:t", ns))
    ppr = para.find("w:pPr", ns)
    if texts or ppr is not None:
        print("XML_P", index, repr(texts), ET.tostring(ppr, encoding="unicode") if ppr is not None else "")
